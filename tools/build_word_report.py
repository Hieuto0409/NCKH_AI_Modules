#!/usr/bin/env python3
r"""
build_word_report.py
--------------------
Tạo báo cáo NCKH cập nhật hoàn chỉnh:
  D:/NCKH/Project_Final/Bao_cao_NCKH_tich_hop_PPG_ECG_SpO2_cap_nhat_25092026.docx

Nội dung cập nhật chính:
  1. ECG AF/non-AF:
     - Đã xác minh thứ tự và đơn vị 9 đặc trưng đầu vào từ dữ liệu huấn luyện (ECG.rar bản _B):
       7 đặc trưng thời gian (giây), pnn50 (%), cv_rr (không thứ nguyên).
       Khớp hoàn toàn với tham số scaler_mean trong model_variables.h.
     - Cửa sổ tính toán 30 giây (15000 mẫu ở 500 Hz), lọc RR trong [0.2, 2.0] giây.
     - Demo trích xuất được đặc trưng nhưng CHƯA CHẠY SUY LUẬN Edge Impulse để trả AF/non-AF.
     - Build firmware PlatformIO chỉ chứng minh biên dịch và liên kết được, KHÔNG chứng minh mô hình đã chạy trên ESP32-S3.
     - Không viết "mô hình ECG chỉ được kích hoạt trong firmware" vì chưa có log suy luận thực tế trên bo.
     - Không viết "xác minh đầy đủ tính đúng đắn của dữ liệu huấn luyện ECG" hay "hoàn thiện toàn bộ hệ thống".
       Diễn đạt đúng: nguyên mẫu phần mềm offline đã kiểm tra một số đường xử lý; cần nối cảm biến, chạy suy luận và thử trên phần cứng.
  2. SpO₂:
     - Giữ nguyên thuật toán Stream100 (MaximCore/ResearchSpO2).
     - Kết quả 99%, 100%, 99% là REPLAY FIXTURE 100 Hz, không phải đo người thật và không chứng minh độ chính xác lâm sàng.
     - Bổ sung 4 vùng phân loại tham khảo lâm sàng có căn cứ:
       * 95–100%: Trong khoảng tham khảo (MedlinePlus / FDA / BTS)
       * 93–94%: Cần chú ý (NHS England COVID Oximetry @home)
       * <= 92%: Cảnh báo SpO₂ thấp (NHS England COVID Oximetry @home)
       * Dữ liệu lỗi / không đạt QC / ngoài [0, 100]: Chưa có kết quả tin cậy
       (Áp dụng người lớn lúc nghỉ, gần mực nước biển, không có mục tiêu SpO2 riêng).
     - Dữ liệu Step 2 (25 Hz) chưa đủ điều kiện kết nối vào SpO2.
  3. Nhịp tim PPG (BPM):
     - Công thức 60000 / median(valid_ppi_ms), tái sử dụng khoảng PPI từ đỉnh PPG 60s.
     - Phân loại nhịp lúc nghỉ theo AHA & NHLBI (<60, 60-100, >100 BPM) dựa trên BPM gốc chưa làm tròn.
     - Nguyên tắc an toàn bối cảnh: Dữ liệu Step 2 mặc định ghi "Chưa đủ bối cảnh để đánh giá theo nhịp lúc nghỉ",
       không tự suy diễn từ SQI hay Stress.
     - Khi QC fail hoặc thiếu đỉnh: trả NOT_READY, không xuất 0 BPM.
"""

import sys
from pathlib import Path
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DOCX = Path(r"D:\NCKH\Project_Final\Bao_cao_NCKH_tich_hop_PPG_ECG_SpO2_cap_nhat_25092026.docx")
OUTPUT_DOCX_PROJECT = Path(__file__).resolve().parent.parent / "Bao_cao_NCKH_tich_hop_PPG_ECG_SpO2_cap_nhat_25092026.docx"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def style_table(table, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for j, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]
            if i == 0:
                set_cell_background(cell, "003366")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
                        r.font.size = Pt(10)
                        r.font.name = "Times New Roman"
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F4F8FA")
                else:
                    set_cell_background(cell, "FFFFFF")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        r.font.name = "Times New Roman"

def create_report():
    doc = docx.Document()

    # Cấu hình lề trang tiêu chuẩn (1 inch = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Tiêu đề báo cáo
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("BÁO CÁO CƠ SỞ KHOA HỌC VÀ KẾT QUẢ KIỂM CHỨNG TÍCH HỢP HỆ THỐNG THEO DÕI SỨC KHỎE")
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(15)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 51, 102)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Phân hệ: Stress PPG · Nhịp tim PPG · ECG AF/non-AF · Thuật toán SpO₂")
    sub_run.font.name = "Times New Roman"
    sub_run.font.size = Pt(12)
    sub_run.font.bold = True
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(70, 70, 70)

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta_p.add_run("Dự án NCKH — Báo cáo kiểm chứng thuật toán offline và biên dịch firmware (Cập nhật 25/09/2026)")
    meta_run.font.name = "Times New Roman"
    meta_run.font.size = Pt(10)
    meta_run.font.italic = True

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 1. TÓM TẮT QUẢN LÝ
    # --------------------------------------------------------------------------
    doc.add_heading("1. Tóm tắt quản lý và trạng thái hệ thống", level=1)
    
    doc.add_paragraph(
        "Báo cáo này tổng hợp kết quả nghiên cứu cơ sở khoa học, chuẩn hóa giao diện dữ liệu và kiểm chứng "
        "tích hợp các đường xử lý thuộc nguyên mẫu phần mềm hệ thống giám sát sức khỏe cá nhân. "
        "Hệ thống kết hợp ba nhánh phân tích sinh lý độc lập: phát hiện trạng thái căng thẳng tâm lý (Stress) qua sóng quang thể tích (PPG), "
        "phát hiện rung nhĩ (AF/non-AF) qua điện tâm đồ (ECG), và ước tính độ bão hòa oxy trong máu (SpO₂). "
        "Bên cạnh đó, nhịp tim ước tính (BPM) được trích xuất trực tiếp từ các khoảng xung PPG hợp lệ."
    )
    doc.add_paragraph(
        "Điểm cập nhật kỹ thuật quan trọng trong phiên bản này gồm: (1) Đã xác minh thứ tự và đơn vị của 9 đặc trưng đầu vào "
        "mô hình ECG Edge Impulse từ tập dữ liệu huấn luyện gốc (ECG.rar bản _B), trong đó 7 đặc trưng thời gian dùng giây (s), "
        "pnn50 dùng %, cv_rr là tỷ số không thứ nguyên; (2) Cửa sổ phân tích ECG được chuẩn hóa về 30 giây (15.000 mẫu ở 500 Hz) "
        "và lọc khoảng RR trong dải [0,2; 2,0] giây, đúng với điều kiện tạo tập huấn luyện; "
        "(3) Thiết lập bộ quy tắc diễn giải lâm sàng tham khảo chính thống cho SpO₂ (theo MedlinePlus & NHS England COVID Oximetry @home) "
        "và cho nhịp tim PPG (theo AHA & NHLBI) dựa trên giá trị chưa làm tròn; "
        "(4) Thực hiện nguyên tắc an toàn bối cảnh (context safety): không tự suy đoán người đo đang ở trạng thái nghỉ ngơi khi chưa có dữ kiện xác nhận."
    )
    doc.add_paragraph(
        "Tất cả kết quả trong báo cáo này được thực nghiệm hoàn toàn trên dữ liệu offline và xác nhận biên dịch firmware PlatformIO. "
        "Nhóm nghiên cứu chưa có phần cứng ESP32-S3 thực tế để nạp mã, chưa thực hiện đo trực tiếp trên người tình nguyện, "
        "và không tự tạo bất kỳ dữ liệu lâm sàng nhân tạo nào."
    )

    doc.add_paragraph("Bảng 1: Trạng thái tổng hợp các chỉ số và mức độ kiểm chứng trong dự án", style="Caption")
    t0 = doc.add_table(rows=5, cols=4)
    headers = ["Nhánh / Chỉ số", "Trạng thái", "Kết quả thực nghiệm", "Mức độ kiểm chứng"]
    for j, h in enumerate(headers):
        t0.cell(0, j).text = h
    
    rows_data = [
        ("Stress PPG", "RESULT_AVAILABLE (✅)", "Baseline (p = 0.198364 trên CSV Step 2)", "Python và C++ khớp tuyệt đối (diff <= 1e-9); test trên WESAD"),
        ("Nhịp tim PPG (BPM)", "RESULT_AVAILABLE (✅)", "88.24 BPM (trung vị PPI = 680 ms, 85 nhịp)", "Thuật toán offline; nhãn tham khảo: Chưa đủ bối cảnh đánh giá nhịp nghỉ"),
        ("ECG AF/non-AF", "NOT_READY (❌)", "9 đặc trưng trích xuất OK (cửa sổ 30s, [0.2, 2.0]s); chưa chạy suy luận EI", "Đơn vị đã xác minh (giây/%/1); firmware build OK chứng minh liên kết, chưa chạy trên bo"),
        ("SpO₂ (Stream100)", "TEST_FIXTURE (🔬)", "Replay fixture 100 Hz: S1=99%, S2=100%, S3=99%", "Replay fixture 100 Hz nội bộ module; không phải đo người thật, chưa chứng minh lâm sàng"),
    ]
    for i, rdata in enumerate(rows_data):
        for j, val in enumerate(rdata):
            t0.cell(i+1, j).text = val
    style_table(t0, [Inches(1.5), Inches(1.6), Inches(2.2), Inches(2.2)])

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 2. KIẾN TRÚC PHẦN MỀM VÀ RANH GIỚI
    # --------------------------------------------------------------------------
    doc.add_heading("2. Kiến trúc phần mềm và ranh giới phần cứng", level=1)
    doc.add_paragraph(
        "Mã nguồn dự án được tổ chức theo kiến trúc module hóa hướng nhúng trong PlatformIO. "
        "Bộ điều phối trung tâm (health_orchestrator) đảm nhiệm việc tiếp nhận dữ liệu hoặc đặc trưng đã qua lọc, "
        "gọi các hàm xử lý tương ứng và trả về cấu trúc kết quả chuẩn hóa (HealthSummaryResult)."
    )
    doc.add_paragraph(
        "Để đảm bảo tính trung thực khoa học, dự án phân định ranh giới nghiêm ngặt giữa ba tầng: "
        "(1) Tầng thu nhận và xử lý tín hiệu phần cứng (thuộc gói Step 2 của nhóm phụ trách phần cứng); "
        "(2) Tầng thuật toán và mô hình AI (bao gồm mã C++ độc lập của Stress PPG, ResearchSpO2 và ECG Edge Impulse); "
        "(3) Tầng kiểm thử offline và cầu nối dữ liệu (các công cụ Python và fixture C++)."
    )

    doc.add_paragraph("Bảng 2: Giao diện dữ liệu và ranh giới kết nối của ba nhánh", style="Caption")
    t1 = doc.add_table(rows=4, cols=3)
    headers1 = ["Nhánh thuật toán", "Yêu cầu giao diện đầu vào", "Tình trạng kết nối thực tế"]
    for j, h in enumerate(headers1):
        t1.cell(0, j).text = h
    rows1 = [
        ("Stress PPG & BPM", "PPG 60 giây; kiểm tra chất lượng SQI; trích xuất 14 đặc trưng HR/PRV", "Demo trích xuất thành công từ ppg_raw(5).csv (25 Hz); C++ binary khớp Python."),
        ("ECG AF/non-AF", "Cửa sổ 30 giây; ECG ~500 Hz; phát hiện đỉnh R; 9 đặc trưng RR", "Trích xuất 9 đặc trưng đúng thứ tự và đơn vị; firmware ESP32-S3 build thành công; chưa chạy suy luận EI trên ESP32-S3."),
        ("SpO₂ (Stream100)", "Dữ liệu thô 2 kênh RED và IR ở 100 Hz có đủ thành phần DC", "Replay fixture 100 Hz PASS; dữ liệu Step 2 (25 Hz, chưa map RED/IR) chưa đủ điều kiện kết nối."),
    ]
    for i, rdata in enumerate(rows1):
        for j, val in enumerate(rdata):
            t1.cell(i+1, j).text = val
    style_table(t1, [Inches(1.8), Inches(2.6), Inches(3.1)])

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 3. CƠ SỞ KHOA HỌC TỪNG NHÁNH
    # --------------------------------------------------------------------------
    doc.add_heading("3. Cơ sở khoa học của các thuật toán tích hợp", level=1)

    # 3.1 Stress PPG
    doc.add_heading("3.1 Mô hình phát hiện Stress qua PPG (WESAD)", level=2)
    doc.add_paragraph(
        "Mô hình Stress PPG được phát triển dựa trên bộ dữ liệu chuẩn WESAD (Wearable Stress and Affect Detection) [1]. "
        "Quy trình xử lý tín hiệu lọc dải thông [0.5, 4.0] Hz để cô lập sóng xung huyết áp thể tích (BVP), "
        "tìm đỉnh tâm thu và tính toán các khoảng cách giữa các đỉnh xung liên tiếp (Peak-to-Peak Interval - PPI). "
        "Từ chuỗi PPI trên cửa sổ trôi 60 giây, 14 đặc trưng thời gian và biến thiên nhịp được trích xuất: "
        "mean_hr_bpm, std_hr_bpm, min_hr_bpm, max_hr_bpm, mean_pp_ms, median_pp_ms, sdnn_ms, rmssd_ms, sdsd_ms, "
        "pnn20_pct, pnn50_pct, cvnn, beat_count, và valid_rr_ratio."
    )
    doc.add_paragraph(
        "Sau khi chuẩn hóa Z-score qua scaler (mean và scale), mô hình hồi quy logistic tính tổ hợp tuyến tính z và xuất xác suất stress: "
        "p = 1 / (1 + exp(-z)). Ngưỡng phân loại là 0.5. "
        "Về mặt thuật ngữ y sinh, biến thiên khoảng xung từ PPG phản ánh biến thiên thể tích xung mạch (Pulse Rate Variability - PRV), "
        "không được tự đồng nhất hoàn toàn với biến thiên nhịp tim (Heart Rate Variability - HRV) bắt nguồn từ phức bộ QRS điện tâm đồ. "
        "Nghiên cứu lâm sàng ghi nhận PRV và HRV có thể có sai lệch nhất định dưới tác động của trương lực vận mạch ngoại vi [2]."
    )

    # 3.2 Nhịp tim PPG
    doc.add_heading("3.2 Thuật toán nhịp tim PPG và ngưỡng tham khảo lâm sàng", level=2)
    doc.add_paragraph(
        "Để tối ưu hóa tài nguyên tính toán và đảm bảo tính nhất quán nội tại của hệ thống, thuật toán nhịp tim (BPM) "
        "tái sử dụng trực tiếp các khoảng PPI hợp lệ đã được trích xuất từ quy trình phát hiện đỉnh PPG 60 giây, "
        "tuyệt đối không chạy thêm một bộ lọc hoặc bộ phát hiện đỉnh thứ hai gây lãng phí bộ nhớ vi điều khiển."
    )
    doc.add_paragraph(
        "Công thức tính nhịp tim phiên đo sử dụng giá trị trung vị (median) của các khoảng PPI sinh lý hợp lệ [333.3 ms – 1500.0 ms]:\n"
        "BPM = 60000 / median(PPI_ms)\n"
        "Phương pháp dùng trung vị mang lại khả năng kháng nhiễu vượt trội so với trung bình cộng, "
        "giúp triệt tiêu ảnh hưởng của các ngoại tâm thu hoặc các xung PPG bị biến dạng cơ học cục bộ [3]. "
        "Khi áp dụng trên dữ liệu mẫu Step 2 (ppg_raw(5).csv), thuật toán phát hiện 88 đỉnh, 85 khoảng nhịp hợp lệ, "
        "trung vị PPI đạt 680.0 ms, tương ứng với nhịp tim ổn định 88.24 BPM."
    )
    doc.add_paragraph(
        "Về diễn giải lâm sàng, theo hướng dẫn của Hiệp hội Tim mạch Hoa Kỳ (AHA) và Viện Tim, Phổi và Máu Quốc gia Hoa Kỳ (NHLBI) [8, 9]: "
        "Khoảng nhịp tim bình thường của người trưởng thành lúc nghỉ (Resting Heart Rate) là từ 60 đến 100 BPM. "
        "Tuy nhiên, việc đánh giá chỉ có giá trị khi biết chắc chắn người đo đang ở trạng thái nghỉ ngơi thể chất và ổn định tinh thần. "
        "Nhịp tim dưới 60 BPM (bradycardia) có thể là sinh lý bình thường ở vận động viên hoặc khi đang ngủ sâu; "
        "trong khi nhịp tim trên 100 BPM (tachycardia) thường xuất hiện tự nhiên khi vận động, sốt, lo âu, hoặc dùng chất kích thích. "
        "Việc phân loại được thực hiện dựa trên giá trị BPM gốc chưa làm tròn (ví dụ 100.004 BPM thuộc phía >100 dù hiển thị làm tròn 100.00 BPM). "
        "Đồng thời, nguyên mẫu áp dụng nguyên tắc an toàn bối cảnh: Nếu không có dữ kiện xác nhận người đo đang nghỉ (is_resting is not True), "
        "hệ thống chỉ hiển thị trị số đo được kèm nhãn 'Chưa đủ bối cảnh để đánh giá theo nhịp lúc nghỉ', "
        "tuyệt đối không tự suy diễn người đo đang nghỉ dựa trên chỉ số SQI hay nhãn Stress."
    )

    # 3.3 ECG AF/non-AF
    doc.add_heading("3.3 Nhánh ECG AF/non-AF và xác minh dữ liệu huấn luyện", level=2)
    doc.add_paragraph(
        "Nhánh phân loại rung nhĩ sử dụng mô hình mạng nơ-ron nhúng được huấn luyện qua nền tảng Edge Impulse (project 1119067). "
        "Qua khảo sát tập tin cấu trúc model_metadata.h dòng 106 và mã nguồn sinh tự động, mô hình thuộc cấu hình Sensor Fusion, "
        "nhận đầu vào đúng 9 đặc trưng HRV theo thứ tự nghiêm ngặt:\n"
        "1. mean_rr: Khoảng RR trung bình\n"
        "2. median_rr: Trung vị khoảng RR\n"
        "3. sdnn: Độ lệch chuẩn các khoảng RR\n"
        "4. rmssd: Căn bậc hai trung bình bình phương các hiệu số RR liên tiếp\n"
        "5. pnn50: Tỷ lệ phần trăm các hiệu số RR vượt quá 50 ms\n"
        "6. cv_rr: Hệ số biến thiên (sdnn / mean_rr)\n"
        "7. iqr_rr: Khoảng tứ phân vị của RR (Q75 - Q25)\n"
        "8. min_rr: Khoảng RR nhỏ nhất\n"
        "9. max_rr: Khoảng RR lớn nhất"
    )
    doc.add_paragraph(
        "Xác minh thứ tự và đơn vị từ dữ liệu huấn luyện gốc (ECG.rar): Khảo sát tập dữ liệu huấn luyện gốc (bản _B - RR đã lọc) "
        "cho thấy 7 đặc trưng thời gian (mean_rr, median_rr, sdnn, rmssd, iqr_rr, min_rr, max_rr) được tính bằng đơn vị GIÂY (s); "
        "pnn50 được tính bằng PHẦN TRĂM (%); cv_rr là tỷ số không thứ nguyên. "
        "Giá trị trung bình của 9 cột này trong tập train khớp chính xác với mảng chuẩn hóa scaler mean trong model_variables.h: "
        "[mean_rr ≈ 0.7669s, median_rr ≈ 0.7634s, sdnn ≈ 0.0895s, rmssd ≈ 0.1247s, pnn50 ≈ 36.014%, cv_rr ≈ 0.1202, "
        "iqr_rr ≈ 0.1007s, min_rr ≈ 0.5874s, max_rr ≈ 0.9810s]. "
        "Do đó, báo cáo này xác nhận thứ tự và đơn vị của 9 đặc trưng đầu vào đã khớp với dữ liệu huấn luyện."
    )
    doc.add_paragraph(
        "Chuẩn hóa pipeline và ranh giới thực thi: Pipeline xử lý ECG demo đã được điều chỉnh về đúng cửa sổ 30 giây (15.000 mẫu ở 500 Hz) "
        "và lọc khoảng RR trong dải [0,2; 2,0] giây, đồng bộ với điều kiện tạo dữ liệu huấn luyện. "
        "Đồng thời, báo cáo ghi nhận sự khác biệt về nguồn gốc: dữ liệu huấn luyện sử dụng vị trí nhịp từ chú giải chuyên gia .qrs ở 250 Hz, "
        "trong khi demo trích xuất đỉnh R bằng thuật toán Pan-Tompkins ở 500 Hz. "
        "Về mặt triển khai, mô hình Edge Impulse đã được biên dịch và liên kết thành công trong firmware PlatformIO. "
        "Tuy nhiên, việc build firmware thành công chỉ xác nhận tính tương thích của API và thư viện, chưa chứng minh mô hình đã chạy trên ESP32-S3 thực tế "
        "do chưa có bo mạch để nạp và thu log suy luận. "
        "Demo offline trích xuất được 9 đặc trưng nhưng chưa chạy suy luận Edge Impulse để trả kết quả AF/non-AF do môi trường PC x86 thiếu runtime TFLite Micro tương thích. "
        "Giữ nguyên trạng thái NOT_READY, kiên quyết không tạo nhãn giả lập hoặc sử dụng quy tắc heuristic tự viết."
    )

    # 3.4 SpO2
    doc.add_heading("3.4 Thuật toán đo SpO₂ và giới hạn kết nối", level=2)
    doc.add_paragraph(
        "Thuật toán SpO₂ tích hợp trong module ResearchSpO2 dựa trên nguyên lý hấp thụ quang phổ của oxyhemoglobin (HbO₂) và deoxyhemoglobin (Hb) "
        "ở hai bước sóng đỏ (RED ~660 nm) và hồng ngoại (IR ~880 nm) [4]. "
        "Tín hiệu PPG của mỗi kênh được tách thành thành phần tĩnh DC và thành phần xung động động mạch AC. "
        "Tỷ số tỷ lệ R = (AC_red / DC_red) / (AC_ir / DC_ir) được ánh xạ qua đường cong hiệu chuẩn thực nghiệm để suy ra tỷ lệ SpO₂. "
        "Lõi xử lý MaximCore thực hiện kiểm tra tưới máu, phát hiện mất tiếp xúc ngón tay và tín hiệu bão hòa ADC."
    )
    doc.add_paragraph(
        "Trong các phép thử replay test độc lập bằng dữ liệu replay/fixture 100 Hz đi kèm module (session_1, session_2, session_3), "
        "thuật toán ghi nhận kết quả 99%, 100%, 99%, đạt trạng thái TEST_FIXTURE. "
        "Cần nhấn mạnh đây là kết quả từ replay fixture 100 Hz nhằm kiểm tra tính đúng đắn của thuật toán nội bộ, "
        "tuyệt đối không phải là phép đo trên người thật và không chứng minh độ chính xác lâm sàng của thiết bị. "
        "Dữ liệu thực tế thu từ phần cứng Step 2 (ppg_raw(5).csv) hiện tại chưa đủ điều kiện kết nối do: "
        "(1) Tần số lấy mẫu bất đồng: Step 2 thu ở 25 Hz, trong khi ResearchSpO2 yêu cầu 100 Hz; "
        "(2) Kênh quang học chưa được định danh: chưa xác định slot nào là RED, slot nào là IR; "
        "(3) Ràng buộc bảo toàn dữ liệu: không tự ý nội suy nhân tạo 25 Hz lên 100 Hz và không đưa tín hiệu đã lọc AC vào hàm pushSpo2."
    )

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 4. DIỄN GIẢI NGƯỠNG LÂM SÀNG
    # --------------------------------------------------------------------------
    doc.add_heading("4. Diễn giải các ngưỡng tham khảo lâm sàng có căn cứ", level=1)
    doc.add_paragraph(
        "Để đảm bảo an toàn và tính chuẩn xác khoa học trong giao diện hiển thị của nguyên mẫu, "
        "hệ thống áp dụng các ngưỡng tham khảo lâm sàng từ các tổ chức y tế quốc tế có uy tín. "
        "Trước khi gắn nhãn, hệ thống bắt buộc kiểm tra giá trị đo là số hữu hạn và nằm trong dải kết quả hợp lệ của module (0 - 100% cho SpO2, > 0 cho BPM). "
        "Mọi nhãn hiển thị đều mang tính chất tham khảo kỹ thuật cho nguyên mẫu nghiên cứu, "
        "không có giá trị chẩn đoán y khoa hay thay thế ý kiến của bác sĩ chuyên khoa."
    )

    doc.add_paragraph("Bảng 3: Khung phân loại SpO₂ tham khảo lâm sàng (MedlinePlus & NHS England)", style="Caption")
    t2 = doc.add_table(rows=5, cols=3)
    headers2 = ["Mức SpO₂ (%)", "Nhãn hiển thị nguyên mẫu", "Căn cứ khoa học và bối cảnh áp dụng"]
    for j, h in enumerate(headers2):
        t2.cell(0, j).text = h
    rows2 = [
        ("95% – 100%", "Trong khoảng tham khảo", "MedlinePlus, FDA và Hiệp hội Lồng ngực Anh (BTS) ghi nhận là dải bình thường ở phần lớn người khỏe mạnh đo lúc nghỉ gần mực nước biển [5, 7, 10]."),
        ("93% – 94%", "Cần chú ý", "Hướng dẫn theo dõi tại nhà COVID Oximetry @home của NHS England: mức cần theo dõi sát, đánh giá lại triệu chứng khó thở và tiếp xúc cảm biến [11]."),
        ("≤ 92%", "Cảnh báo SpO₂ thấp", "NHS England: ngưỡng cảnh báo đỏ cần hỗ trợ y tế khẩn cấp trong chương trình theo dõi oxy xung tại nhà [11]. (WHO xem dưới 90% là khẩn cấp trong gây mê [6])."),
        ("Dữ liệu lỗi / Mất tiếp xúc / Ngoài [0, 100]", "Chưa có kết quả tin cậy", "Tín hiệu yếu, cử động quá mức, bão hòa ADC, ngoài dải hoặc không đạt SQI; trả valid=False, percent=None, không xuất số đo tin cậy [4, 5]."),
    ]
    for i, rdata in enumerate(rows2):
        for j, val in enumerate(rdata):
            t2.cell(i+1, j).text = val
    style_table(t2, [Inches(1.8), Inches(2.2), Inches(3.5)])

    doc.add_paragraph()

    doc.add_paragraph("Bảng 4: Khung phân loại nhịp tim lúc nghỉ tham khảo lâm sàng (AHA & NHLBI)", style="Caption")
    t3 = doc.add_table(rows=5, cols=3)
    headers3 = ["Mức nhịp tim (BPM)", "Nhãn hiển thị nguyên mẫu", "Căn cứ khoa học và lưu ý bối cảnh"]
    for j, h in enumerate(headers3):
        t3.cell(0, j).text = h
    rows3 = [
        ("< 60 BPM", "Thấp hơn khoảng tham khảo lúc nghỉ", "AHA & NHLBI: nhịp chậm (bradycardia). Đánh giá dựa trên BPM gốc chưa làm tròn; chỉ áp dụng khi xác nhận người đo đang nghỉ (is_resting=True) [8, 9]."),
        ("60 – 100 BPM", "Trong khoảng tham khảo lúc nghỉ", "AHA & NHLBI: dải nhịp tim lúc nghỉ tiêu chuẩn của người trưởng thành khỏe mạnh (is_resting=True); ví dụ 100.000 BPM thuộc dải này [8, 9]."),
        ("> 100 BPM", "Cao hơn khoảng tham khảo lúc nghỉ", "AHA & NHLBI: nhịp nhanh (tachycardia). Đánh giá dựa trên BPM gốc chưa làm tròn; ví dụ 100.004 BPM thuộc dải này dù hiển thị 100.00 BPM [8, 9]."),
        ("Bối cảnh chưa xác định (Mặc định)", "Chưa đủ bối cảnh để đánh giá theo nhịp lúc nghỉ", "Áp dụng khi hệ thống chưa có dữ kiện xác nhận người đo đang nghỉ ngơi; hiển thị trị số BPM thực nhưng không phân loại dải sinh lý [8]."),
    ]
    for i, rdata in enumerate(rows3):
        for j, val in enumerate(rdata):
            t3.cell(i+1, j).text = val
    style_table(t3, [Inches(2.0), Inches(2.3), Inches(3.2)])

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 5. KẾT QUẢ KIỂM TRA THỰC NGHIỆM
    # --------------------------------------------------------------------------
    doc.add_heading("5. Kết quả kiểm tra thực nghiệm và đối chiếu chéo", level=1)
    doc.add_paragraph(
        "Toàn bộ hệ thống mã nguồn đã trải qua bộ kiểm thử nghiêm ngặt bao gồm unit test, kiểm tra đối chiếu chéo (cross-validation) "
        "giữa Python và C++, kiểm tra dữ liệu replay/fixture 100 Hz đi kèm module và biên dịch firmware vi điều khiển."
    )

    doc.add_paragraph("Bảng 5: Tổng hợp kết quả thực nghiệm bộ kiểm thử của dự án", style="Caption")
    t4 = doc.add_table(rows=9, cols=3)
    headers4 = ["Lệnh kiểm tra / Tập tin test", "Kết quả quan sát được", "Ý nghĩa kỹ thuật"]
    for j, h in enumerate(headers4):
        t4.cell(0, j).text = h
    rows4 = [
        ("python test/test_interpretation_rules.py", "10/10 ca đạt (Ran 10 tests - OK)", "Xác nhận đúng logic phân vùng SpO₂ (biên 92, 93, 94, 95, 100; lỗi NaN, Inf, -1, 105), BPM unrounded (100.004 -> >100, 59.996 -> <60), và luồng đầy đủ PPI -> BPM -> phân loại."),
        ("python test/test_ppg_heart_rate.py", "6/6 ca đạt (PASS ✅)", "Xác nhận công thức 60000/median(PPI); PPI 800 ms -> 75 BPM; trung vị lọc nhiễu; bảo toàn BPM gốc 100.004/59.996 đến lúc phân loại; từ chối trả 0 BPM khi SQI fail."),
        ("python tools/run_offline_demo.py", "Exit code 0; JSON xuất thành công", "Chạy hoàn chỉnh luồng demo offline 3 nhánh; hiển thị bảng tổng hợp trung thực; thời gian chạy ~4.6s."),
        ("python test/verify_stress.py", "Deployment: 58/60 (96.67%); Eval: 56/60 (93.33%)", "Kiểm thử hồi quy trên toàn bộ 60 cửa sổ WESAD (S13, S16); 0 ca False Negative."),
        (".pio/stress_test.exe (C++ binary)", "6/6 fixture WESAD đạt (diff <= 1e-9)", "Đối chiếu độc lập giữa Python và C++ binary; xác suất tính toán trùng khớp hoàn toàn."),
        ("SPO2_Module/spo2_test.exe", "Replay session 1, 2, 3: 99%, 100%, 99% (PASS)", "Kiểm thử thuật toán ResearchSpO2::Stream100 trên dữ liệu replay/fixture 100 Hz đi kèm module; không phải kết quả đo người thật."),
        ("pio run -e esp32-s3-devkitc-1", "BUILD SUCCESS (118.6s; Flash 16.4%, RAM 8.9%)", "Chứng minh mã nguồn biên dịch và liên kết thành công trên môi trường ESP32-S3; chưa chứng minh đã chạy trên bo thực."),
        ("pio run -e test_fixture", "BUILD SUCCESS (101.8s; Flash 16.6%, RAM 8.9%)", "Chứng minh cấu hình lib_ignore hoạt động đúng, không còn xung đột thư viện trùng lặp."),
    ]
    for i, rdata in enumerate(rows4):
        for j, val in enumerate(rdata):
            t4.cell(i+1, j).text = val
    style_table(t4, [Inches(2.5), Inches(2.5), Inches(2.5)])

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 6. GIỚI HẠN KỸ THUẬT VÀ BÀN GIAO
    # --------------------------------------------------------------------------
    doc.add_heading("6. Giới hạn kỹ thuật và kế hoạch bàn giao phần cứng", level=1)
    doc.add_paragraph(
        "Nhóm nghiên cứu bảo lưu các giới hạn thực tế để đảm bảo tính khách quan và đạo đức khoa học:\n"
        "1. Trạng thái mô hình ECG: Đã xác minh thứ tự và đơn vị 9 đặc trưng đầu vào (giây, %, 1) từ dữ liệu huấn luyện ECG.rar. "
        "Demo trích xuất được đặc trưng nhưng chưa chạy suy luận Edge Impulse để trả AF/non-AF. "
        "Việc build firmware PlatformIO thành công chỉ chứng minh mã nguồn biên dịch và liên kết được, "
        "chưa chứng minh mô hình đã chạy trên vi điều khiển ESP32-S3 thực tế vì chưa có phần cứng để nạp firmware và ghi nhận log suy luận.\n"
        "2. Đồng bộ tần số lấy mẫu PPG: Cần nhóm phần cứng cấu hình thanh ghi cảm biến MAX30102 phát xung ở tần số 100 Hz "
        "để đáp ứng đồng thời cả thuật toán SpO₂ (Stream100) và thuật toán nhịp tim/PRV.\n"
        "3. Định danh kênh quang học: Cần tài liệu xác nhận rõ ràng chân LED nào ứng với RED (660 nm) và IR (880 nm) trên module phần cứng thực tế.\n"
        "4. Kiểm chứng lâm sàng: Toàn bộ kết quả hiện có là kiểm chứng thuật toán offline trên tập dữ liệu mở (WESAD, PhysioNet, dữ liệu replay/fixture 100 Hz đi kèm module). "
        "Đề tài chưa tiến hành thử nghiệm lâm sàng trên người bệnh thật sự và không đưa ra tuyên bố về độ chính xác chẩn đoán y khoa."
    )
    doc.add_paragraph(
        "Kế hoạch bàn giao cho giai đoạn kế tiếp:\n"
        "• Bàn giao gói mã nguồn PlatformIO đã tối ưu, cấu hình lib_ignore bỏ qua các thư mục sao lưu trùng lặp, "
        "biên dịch thành công cả hai môi trường.\n"
        "• Bàn giao công cụ kiểm thử tự động một lệnh `python tools/run_offline_demo.py` và bộ unit test `test_interpretation_rules.py` "
        "để nhóm phần cứng có thể tự động kiểm tra tính toàn vẹn của thuật toán ngay khi có bo mạch ESP32-S3.\n"
        "• Cung cấp tài liệu giao diện API của health_orchestrator để nối luồng đọc dữ liệu từ I2C và ADC thực tế."
    )

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # 7. KẾT LUẬN
    # --------------------------------------------------------------------------
    doc.add_heading("7. Kết luận", level=1)
    doc.add_paragraph(
        "Bản cập nhật hiện tại của dự án AI_Moudel_Summary là nguyên mẫu phần mềm offline đã kiểm tra thành công một số đường xử lý quan trọng: "
        "xác minh thứ tự và đơn vị 9 đặc trưng ECG từ dữ liệu huấn luyện, chuẩn hóa công thức nhịp tim BPM từ PPI và thiết lập khung phân loại tham khảo lâm sàng minh bạch cho SpO₂ và nhịp tim. "
        "Tuy nhiên, để hệ thống hoạt động thực sự trên thực địa, nhóm nghiên cứu cần tiếp tục kết nối luồng cảm biến thực tế (I2C MAX30102, ADC AD8232), "
        "thực thi suy luận mô hình và thử nghiệm đo đạc trực tiếp trên phần cứng ESP32-S3."
    )

    doc.add_paragraph()

    # --------------------------------------------------------------------------
    # TÀI LIỆU THAM KHẢO
    # --------------------------------------------------------------------------
    doc.add_heading("Tài liệu tham khảo", level=1)
    refs = [
        "[1] Schmidt, P., et al. Introducing WESAD, a Multimodal Dataset for Wearable Stress and Affect Detection. ICMI 2018. https://archive.ics.uci.edu/dataset/465/wesad+wearable+stress+and+affect+detection",
        "[2] Yuda, E., et al. Pulse rate variability is not the same as heart rate variability. Frontiers in Physiology, 2025. https://www.frontiersin.org/journals/physiology/articles/10.3389/fphys.2025.1566860/full",
        "[3] Analog Devices. MAXREFDES117#: Heart-Rate and Pulse-Oximetry Monitor. Reference Design. https://www.analog.com/en/resources/reference-designs/maxrefdes117.html",
        "[4] Analog Devices. Guidelines for SpO2 Measurement Using the Maxim MAX30101/MAX30102. Application Note. https://www.analog.com/en/resources/app-notes/guidelines-for-spo2-measurement.html",
        "[5] U.S. Food and Drug Administration (FDA). Pulse Oximeter Accuracy and Limitations: FDA Safety Communication. https://www.fda.gov/medical-devices/safety-communications/pulse-oximeter-accuracy-and-limitations-fda-safety-communication",
        "[6] World Health Organization (WHO). Pulse Oximetry Training Manual. Geneva: World Health Organization; 2011. https://www.who.int/publications/i/item/9789241501132",
        "[7] O'Driscoll, B. R., et al. British Thoracic Society Guideline for oxygen use in adults in healthcare and emergency settings. BMJ Open Respiratory Research, 2017. https://www.brit-thoracic.org.uk/quality-improvement/guidelines/emergency-oxygen/",
        "[8] American Heart Association (AHA). All About Heart Rate (Pulse). https://www.heart.org/en/health-topics/high-blood-pressure/the-facts-about-high-blood-pressure/all-about-heart-rate-pulse",
        "[9] National Heart, Lung, and Blood Institute (NHLBI). Types of Arrhythmias. https://www.nhlbi.nih.gov/health/arrhythmias/types",
        "[10] MedlinePlus. Pulse Oximetry. U.S. National Library of Medicine. https://medlineplus.gov/lab-tests/pulse-oximetry/",
        "[11] NHS England. COVID Oximetry @home: standard operating procedure. https://www.england.nhs.uk/coronavirus/publication/covid-oximetry-at-home-sop/",
    ]
    for ref in refs:
        p_ref = doc.add_paragraph(ref)
        p_ref.paragraph_format.left_indent = Inches(0.3)
        p_ref.paragraph_format.first_line_indent = Inches(-0.3)
        for r in p_ref.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)

    doc.save(str(OUTPUT_DOCX))
    doc.save(str(OUTPUT_DOCX_PROJECT))
    print(f"Báo cáo Word đã được tạo thành công tại:")
    print(f"  [1] {OUTPUT_DOCX}")
    print(f"  [2] {OUTPUT_DOCX_PROJECT}")

if __name__ == "__main__":
    create_report()
